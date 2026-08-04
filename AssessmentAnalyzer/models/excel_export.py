"""
excel_export.py
Builds a formatted Excel workbook (in-memory) summarizing a survey/assessment
analysis: reliability stats, item difficulty/discrimination, and cluster profile.
Returned as bytes so the Flask route can stream it as a download.
"""

import io
import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils.dataframe import dataframe_to_rows

HEADER_FILL = PatternFill(start_color="2E5E4E", end_color="2E5E4E", fill_type="solid")
HEADER_FONT = Font(color="FFFFFF", bold=True)


def _style_header(ws, ncols):
    for col in range(1, ncols + 1):
        c = ws.cell(row=1, column=col)
        c.fill = HEADER_FILL
        c.font = HEADER_FONT
        c.alignment = Alignment(horizontal="center")


def _autofit(ws, ncols):
    for col in range(1, ncols + 1):
        letter = ws.cell(row=1, column=col).column_letter
        max_len = max(len(str(ws.cell(row=r, column=col).value)) for r in range(1, ws.max_row + 1))
        ws.column_dimensions[letter].width = min(max_len + 3, 35)


def build_report(item_df: pd.DataFrame, report: dict, cluster_info: dict | None = None) -> bytes:
    wb = Workbook()

    # --- Sheet 1: Summary KPIs ---
    ws1 = wb.active
    ws1.title = "Summary"
    ws1.append(["Metric", "Value"])
    ws1.append(["Respondents", report["n_respondents"]])
    ws1.append(["Items", report["n_items"]])
    ws1.append(["Cronbach's Alpha", report["cronbach_alpha"]])
    _style_header(ws1, 2)
    _autofit(ws1, 2)

    # --- Sheet 2: Item statistics ---
    ws2 = wb.create_sheet("Item Statistics")
    rows = []
    for item in item_df.columns:
        rows.append({
            "item": item,
            "difficulty": report["item_difficulty"].get(item),
            "discrimination": report["item_discrimination"].get(item),
            "mean": report["descriptives"].get(item, {}).get("mean"),
            "std": report["descriptives"].get(item, {}).get("std"),
        })
    item_stats_df = pd.DataFrame(rows)
    for row in dataframe_to_rows(item_stats_df, index=False, header=True):
        ws2.append(row)
    _style_header(ws2, item_stats_df.shape[1])
    _autofit(ws2, item_stats_df.shape[1])

    # --- Sheet 3: Raw data ---
    ws3 = wb.create_sheet("Raw Responses")
    for row in dataframe_to_rows(item_df, index=False, header=True):
        ws3.append(row)
    _style_header(ws3, item_df.shape[1])
    _autofit(ws3, item_df.shape[1])

    # --- Sheet 4: Cluster profile (optional) ---
    if cluster_info:
        ws4 = wb.create_sheet("Cluster Profile")
        profile_df = pd.DataFrame(cluster_info["cluster_profile"]).T
        profile_df.index.name = "cluster"
        profile_df.reset_index(inplace=True)
        for row in dataframe_to_rows(profile_df, index=False, header=True):
            ws4.append(row)
        _style_header(ws4, profile_df.shape[1])
        _autofit(ws4, profile_df.shape[1])

    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    return buf.read()
